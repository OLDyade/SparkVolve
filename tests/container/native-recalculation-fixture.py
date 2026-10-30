import json
import sys
from datetime import datetime
from pathlib import Path
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

root = Path(sys.argv[2])
formulas = {
    'A1': '=ROUND(2.5,0)', 'A2': '=IF(1=1,7,1/0)',
    'A3': "=SUM('Other'!A1:A2)", 'A4': '=2^3^2',
    'A5': '="#literal"', 'A6': '=1/0', 'A7': '=DATE(2020,2,3)',
}
if sys.argv[1] == 'create':
    w = Workbook()
    s = w.active
    s.title = 'Model'
    o = w.create_sheet('Other')
    o['A1'], o['A2'] = 2, 3
    for address, value in formulas.items(): s[address] = value
    s['A7'].number_format = 'yyyy-mm-dd'
    s['D1'].comment = Comment('An intentionally empty observation with a comment', 'contract')
    w.save(root / 'source.xlsx')
    special = Workbook()
    sheet = special.active
    sheet.title = 'Special'
    sheet['A1'], sheet['A2'], sheet['B1'], sheet['B2'] = 2, 3, 4, 5
    sheet['C1'] = ArrayFormula(ref='C1:C2', text='=A1:A2*B1:B2')
    special.save(root / 'array.xlsx')
    sheet['E1'] = DataTableFormula(ref='E1:F2', dt2D=True, r1='A1', r2='B1', ca=True)
    special.save(root / 'special.xlsx')
elif sys.argv[1] == 'verify-array':
    f = load_workbook(root / 'array-result.xlsx', data_only=False)
    v = load_workbook(root / 'array-result.xlsx', data_only=True)
    assert isinstance(f['Special']['C1'].value, ArrayFormula)
    assert f['Special']['C1'].value.ref == 'C1:C2'
    assert [v['Special'][cell].value for cell in ['C1', 'C2']] == [8, 15]
    print(json.dumps({'arrayResults': [8, 15], 'range': 'C1:C2'}))
elif sys.argv[1] == 'verify':
    f = load_workbook(root / 'result.xlsx', data_only=False)
    v = load_workbook(root / 'result.xlsx', data_only=True)
    assert all(f['Model'][address].data_type == 'f' for address in formulas), 'recalculated delivery lost formulas'
    assert [v['Model'][f'A{i}'].value for i in range(1, 7)] == [3, 7, 5, 64, '#literal', '#DIV/0!']
    assert v['Model']['A5'].data_type == 's'
    assert v['Model']['A6'].data_type == 'e'
    assert v['Model']['A7'].value == datetime(2020, 2, 3)
    assert f['Other']['A1'].value == 2 and f['Other']['A2'].value == 3
    assert f['Model']['D1'].value is None and f['Model']['D1'].comment.text.startswith('An intentionally empty')
    print(json.dumps({'formulaCount': len(formulas), 'cachedResults': 'matched', 'sourceCells': 'preserved'}))
else:
    raise ValueError('Unknown fixture operation')
