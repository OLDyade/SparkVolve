import json
import sys
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple

args = json.load(sys.stdin)
if args['operation'] == 'inspect':
    workbook = load_workbook(args['file'])
    print(json.dumps({sheet.title: [{'cell': cell.coordinate, 'value': cell.value, 'comment': cell.comment.text if cell.comment else None}
                                  for row in sheet for cell in row if cell.value is not None or cell.comment]
                      for sheet in workbook}, default=str))
elif args['operation'] == 'write':
    target = Path(args['destination']).resolve()
    if not target.is_relative_to('/workspace'):
        raise ValueError('Destination must be inside /workspace')
    workbook = load_workbook(args['source'])
    seen = set()
    for edit in args['edits']:
        position = (edit['sheet'], coordinate_to_tuple(edit['cell']))
        if position in seen:
            raise ValueError('Duplicate target')
        seen.add(position)
        workbook[edit['sheet']][edit['cell']] = edit['value']
    target.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(target)
    print(json.dumps({'saved': str(target), 'submitted_cells': len(seen), 'formula_values_recalculated': False}))
else:
    raise ValueError('Unknown workbook operation')
