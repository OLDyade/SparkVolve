"""Create openly specified business fixtures. No historical benchmark materials are read."""
import json
import sys
import shutil
from pathlib import Path
from openpyxl import Workbook
from openpyxl.comments import Comment

root = Path(sys.argv[1])
root.mkdir()
tasks = []
for name, years, revenues, margins, tax in [
    ('operating-plan', [2024, 2025, 2026], [1000, 1200, 1500], [.30, .32, .35], .25),
    ('regional-plan', [2025, 2026, 2027], [850, 1020, 1173], [.22, .27, .29], .20),
    ('stress-plan', [2023, 2024, 2025], [700, 630, 756], [.40, .38, .42], .30),
]:
    public = root / 'inputs' / name
    private = root / 'evaluation' / name
    public.mkdir(parents=True); private.mkdir(parents=True)
    wb = Workbook(); ws = wb.active; ws.title = 'Plan'
    for row in [['Metric', *years], ['Revenue', *revenues], ['Gross margin', *margins], ['Gross profit'], ['Operating cost', 80, 90, 100], ['Operating profit'], ['Tax rate', tax, tax, tax], ['Tax expense'], ['Net profit'], ['Revenue growth']]:
        ws.append(row)
    ws['A3'].comment = Comment('Gross margin is gross profit divided by revenue, not cost divided by revenue.', 'Public task definition')
    ws['A10'].comment = Comment('The first year has no prior observation: leave its growth blank. Compute the other years from their respective prior-year revenues.', 'Public task definition')
    wb.save(public / 'source.xlsx')
    expected = []
    for i, col in enumerate('BCD'):
        gross = revenues[i] * margins[i]; operating = gross - [80,90,100][i]
        for row, value in [(4,gross),(6,operating),(8,operating*tax),(9,operating*(1-tax))]:
            expected.append({'sheet':'Plan','cell':f'{col}{row}','value':value})
        if i: expected.append({'sheet':'Plan','cell':f'{col}10','value':revenues[i]/revenues[i-1]-1})
    (private/'expected.json').write_text(json.dumps(expected))
    shutil.copyfile(Path(__file__).with_name('grade.py'), private/'grade.py')
    tasks.append({'id':name,'input':f'inputs/{name}','evaluation':f'evaluation/{name}', 'prompt':'Complete /input/source.xlsx for every displayed year using the original labels and comments. Fill gross profit, operating profit, tax expense, net profit, and revenue growth with Excel formulas. Keep source inputs and the first-year unavailable growth unchanged. Deliver /workspace/result.xlsx. Formulas will be recalculated independently. Python openpyxl and LibreOffice are installed.'})
(root/'tasks.json').write_text(json.dumps(tasks,indent=2))
