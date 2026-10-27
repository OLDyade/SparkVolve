"""Independent synthetic API fixture, never an answer generator or workbook repairer."""
import json
import sys
from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font

operation, source = sys.argv[1:3]
if operation == 'create':
    book = Workbook()
    sheet = book.active
    sheet.title = 'Contract'
    sheet['A1'], sheet['A2'] = 2, 3
    sheet['C3'], sheet['C4'] = 99, 'explicitly cleared'
    sheet['D5'] = 'preserved label'
    sheet['D5'].comment = Comment('original comment', 'fixture')
    sheet['D5'].font = Font(bold=True, color='FF112233')
    other = book.create_sheet('Other')
    other['A1'] = '=Contract!A1'
    book.save(source)
elif operation == 'verify':
    original = load_workbook(source)
    actual = load_workbook(sys.argv[3])
    edits = json.loads(sys.argv[4])
    targets = {(edit['sheet'], edit['cell']) for edit in edits}
    assert actual.sheetnames == original.sheetnames, actual.sheetnames
    for edit in edits:
        value = actual[edit['sheet']][edit['cell']].value
        assert value == edit['value'], (edit, value)
    for sheet in original:
        for row in sheet:
            for cell in row:
                if (sheet.title, cell.coordinate) in targets:
                    continue
                saved = actual[sheet.title][cell.coordinate]
                assert saved.value == cell.value, (sheet.title, cell.coordinate, saved.value)
                assert saved._style == cell._style, (sheet.title, cell.coordinate, 'style')
                assert saved.comment == cell.comment, (sheet.title, cell.coordinate, 'comment')
else:
    raise ValueError(f'Unknown fixture operation: {operation}')
