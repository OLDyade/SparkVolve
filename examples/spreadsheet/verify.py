"""Verify the actual spreadsheet image before any paid model task; no benchmark data."""
import json
import platform
import runpy
import subprocess
from pathlib import Path
import openpyxl

book = openpyxl.Workbook()
sheet = book.active
sheet.title = 'Calculation'
sheet['A1'] = 7
sheet['B1'] = 'Original source'
book.save('/input/source.xlsx')
sheet['A2'] = '=A1+5'
book.save('/delivery/result.xlsx')
runpy.run_path('/evaluation/grade.py', run_name='__main__')
grade = json.loads(Path('/workspace/result.json').read_text())
assert grade['score'] == 1, grade
assert grade['metrics']['preserved'] == 1, grade
assert openpyxl.load_workbook('/delivery/result.xlsx', data_only=False)['Calculation']['A2'].value == '=A1+5'
assert openpyxl.load_workbook('/workspace/recalculated/result.xlsx', data_only=True)['Calculation']['A2'].value == 12
Path('/workspace/preflight.json').write_text(json.dumps({
    'passed': True,
    'python': platform.python_version(),
    'openpyxl': openpyxl.__version__,
    'libreoffice': subprocess.run(['libreoffice', '--version'], check=True, capture_output=True, text=True).stdout.strip(),
    'grade': grade,
    'boundary': 'Synthetic dependency and independent recalculation verification, no model calls or benchmark score.',
}))
