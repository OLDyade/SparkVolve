"""Independent synthetic business acceptance. This is not SpreadsheetBench's evaluator."""
import json
import math
import subprocess
from pathlib import Path
from openpyxl import load_workbook

delivery = Path('/delivery/result.xlsx')
expected = json.loads(Path('/evaluation/expected.json').read_text())
if not delivery.is_file():
    result = {'score': 0, 'metrics': {'preserved': 0}, 'feedback': 'Required result.xlsx was not delivered'}
else:
    output = Path('/workspace/recalculated')
    output.mkdir()
    subprocess.run(['libreoffice', '-env:UserInstallation=file:///tmp/lo-grade', '--headless', '--convert-to', 'xlsx', '--outdir', str(output), str(delivery)], check=True, capture_output=True, timeout=60)
    actual = load_workbook(output / 'result.xlsx', data_only=True)
    source = load_workbook('/input/source.xlsx', data_only=False)
    submitted = load_workbook(delivery, data_only=False)
    checks = []
    target_keys = {(item['sheet'], item['cell']) for item in expected}
    for item in expected:
        value = actual[item['sheet']][item['cell']].value
        passed = submitted[item['sheet']][item['cell']].data_type == 'f' and isinstance(value, (int, float)) and not isinstance(value, bool) and math.isclose(value, item['value'], rel_tol=1e-9, abs_tol=1e-9)
        checks.append({'sheet': item['sheet'], 'cell': item['cell'], 'passed': passed, 'actual': value})
    unchanged = [submitted[sheet.title][cell.coordinate].value == cell.value for sheet in source for row in sheet for cell in row if (sheet.title, cell.coordinate) not in target_keys]
    result = {'score': sum(item['passed'] for item in checks) / len(checks), 'metrics': {'preserved': sum(unchanged) / len(unchanged)}, 'feedback': json.dumps(checks)}
Path('/workspace/result.json').write_text(json.dumps(result))
