"""XLSX workbook operations for the spreadsheet bundle.

Calculation is delegated to the installed spreadsheet engine (LibreOffice in
headless mode with an isolated temporary user profile). This module does not
interpret, patch or guess formula results.

Request (stdin JSON): {"operation": "inspect" | "write" | "recalculate", ...}
Response (stdout JSON): one JSON document; engine failures raise and are
reported by the caller through a non-zero exit code and stderr text.
"""
import json
import os
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import date, datetime, time as datetime_time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.utils.cell import coordinate_to_tuple
from openpyxl.worksheet.formula import ArrayFormula, DataTableFormula

WORKSPACE_ROOT = '/workspace'
ENGINE_BINARY = 'libreoffice'


def emit(payload):
    sys.stdout.write(json.dumps(payload, default=json_value, allow_nan=False))
    sys.stdout.write('\n')


def json_value(value):
    if isinstance(value, (date, datetime, datetime_time)):
        return value.isoformat()
    raise TypeError('Unsupported workbook value type: %s' % type(value).__name__)


def resolve_inside_workspace(raw, label):
    resolved = Path(raw).resolve()
    if not resolved.is_relative_to(Path(WORKSPACE_ROOT).resolve()):
        raise ValueError('%s must be inside %s' % (label, WORKSPACE_ROOT))
    return resolved


def normalize_value(value):
    """Model-facing representation: integral floats -> int, nothing else changed."""
    if isinstance(value, bool):
        return value
    if isinstance(value, float) and value.is_integer() and abs(value) < 2 ** 53:
        return int(value)
    return value


def is_formula(value):
    return isinstance(value, (ArrayFormula, DataTableFormula)) or (
        isinstance(value, str) and value.startswith('=')
    )


def formula_record(value):
    """Expose loaded OOXML formula text and attributes without interpreting them."""
    if isinstance(value, ArrayFormula):
        return {'formula': value.text, 'formula_attributes': dict(value)}
    if isinstance(value, DataTableFormula):
        # Data tables store input references and flags, not a scalar formula string.
        return {'formula': None, 'formula_attributes': dict(value)}
    if is_formula(value):
        return {'formula': value, 'formula_attributes': {'t': 'normal'}}
    raise TypeError('Expected a stored formula')


def operation_inspect(args):
    source = Path(args['file']).resolve()
    book = load_workbook(source)
    only = args.get('sheets')
    if only:
        missing = [name for name in only if name not in book.sheetnames]
        if missing:
            raise ValueError('Unknown sheet(s): %s' % ', '.join(missing))
    result = {}
    for sheet in book:
        if only and sheet.title not in only:
            continue
        cells = {}
        for row in sheet.iter_rows():
            for cell in row:
                if cell.value is None and cell.comment is None:
                    continue
                entry = {}
                if is_formula(cell.value):
                    entry.update(formula_record(cell.value))
                else:
                    entry['value'] = normalize_value(cell.value)
                if cell.comment is not None:
                    entry['comment'] = cell.comment.text
                cells[cell.coordinate] = entry
        result[sheet.title] = {
            'state': sheet.sheet_state,
            'dimensions': sheet.dimensions,
            'merged': [str(rng) for rng in sheet.merged_cells.ranges],
            'cells': cells,
        }
    return {
        'sheetnames': list(book.sheetnames),
        'sheets': result,
        'calculation_engine_invoked': False,
        'stored_values_only': True,
        'note': 'Reads stored cell content: literal values, formula text, formula attributes and comments. '
                'Stored cached results of formulas are not refreshed here; use the '
                'recalculate operation for engine results.',
    }


def coerce_edit_value(value):
    if isinstance(value, (int, float, str)) or value is None:
        return value
    raise ValueError('Cell value must be a string, number or null')


def operation_write(args):
    destination = resolve_inside_workspace(args['destination'], 'destination')
    source = Path(args['source']).resolve()
    if destination == source:
        raise ValueError('Destination must differ from source')
    book = load_workbook(source)
    seen = set()
    for edit in args['edits']:
        sheet = edit['sheet']
        if sheet not in book.sheetnames:
            raise ValueError('Unknown sheet: %s' % sheet)
        try:
            position = coordinate_to_tuple(edit['cell'])
        except (TypeError, ValueError):
            raise ValueError('Invalid cell reference: %r' % (edit['cell'],))
        if (sheet, position) in seen:
            raise ValueError('Duplicate target %s!%s' % (sheet, edit['cell']))
        seen.add((sheet, position))
        book[sheet][edit['cell']].value = coerce_edit_value(edit['value'])
    destination.parent.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix='.workbook-write-', dir=destination.parent) as tmp:
        staged = Path(tmp) / 'result.xlsx'
        book.save(staged)
        os.replace(staged, destination)
    return {
        'saved': str(destination),
        'submitted_cells': len(seen),
        'calculation_engine_invoked': False,
        'note': 'Formulas were written exactly as supplied. Use the recalculate '
                'operation to observe engine results.',
    }


def find_engine():
    engine = shutil.which(ENGINE_BINARY)
    if engine is None:
        raise RuntimeError('Required LibreOffice executable is not installed')
    return engine


def convert_with_engine(source, out_dir, profile_dir):
    """Run the engine headless on a private copy with an isolated user profile."""
    engine = find_engine()
    profile_dir.mkdir(parents=True, exist_ok=True)
    out_dir.mkdir(parents=True, exist_ok=True)
    command = [
        engine,
        '-env:UserInstallation={0}'.format(profile_dir.resolve().as_uri()),
        '--headless',
        '--norestore',
        '--invisible',
        '--nodefault',
        '--nofirststartwizard',
        '--convert-to',
        'xlsx:Calc MS Excel 2007 XML',
        '--outdir',
        str(out_dir),
        str(source),
    ]
    env = dict(os.environ)
    env['HOME'] = str(profile_dir)
    env['TMPDIR'] = str(profile_dir)
    started = time.time()
    try:
        completed = subprocess.run(
            command,
            capture_output=True,
            text=True,
            timeout=600,
            env=env,
        )
    except subprocess.TimeoutExpired as error:
        raise RuntimeError('Spreadsheet calculation engine timed out: %s' % error)
    detail = {
        'engine': engine,
        'exit_code': completed.returncode,
        'elapsed_seconds': round(time.time() - started, 3),
        'stdout': completed.stdout.strip()[-2000:],
        'stderr': completed.stderr.strip()[-2000:],
    }
    converted = out_dir / (source.stem + '.xlsx')
    detail['output_present'] = converted.is_file()
    if completed.returncode != 0 or not converted.is_file():
        raise RuntimeError('Spreadsheet calculation engine failed: %s' % json.dumps(detail))
    return converted, detail


def operation_recalculate(args):
    source = Path(args['source']).resolve()
    if not source.is_file():
        raise ValueError('Source workbook does not exist: %s' % source)
    destination = None
    if args.get('destination'):
        destination = resolve_inside_workspace(args['destination'], 'destination')
        if destination == source:
            raise ValueError('Destination must differ from source')
    requested = args.get('sheets')
    with tempfile.TemporaryDirectory(prefix='xl-engine-') as tmp:
        tmp_path = Path(tmp)
        isolated = tmp_path / 'source.xlsx'
        shutil.copyfile(source, isolated)
        converted, detail = convert_with_engine(
            isolated, tmp_path / 'out', tmp_path / 'profile'
        )
        original_book = load_workbook(source)
        # data_only reads the cached results the engine stored beside each formula.
        calculated_book = load_workbook(converted, data_only=True)
        if requested:
            missing = [name for name in requested if name not in original_book.sheetnames]
            if missing:
                raise ValueError('Unknown sheet(s): %s' % ', '.join(missing))
        rows = []
        errors = 0
        for sheet in original_book:
            if requested and sheet.title not in requested:
                continue
            counterpart = calculated_book[sheet.title]
            for row in sheet.iter_rows():
                for cell in row:
                    if not is_formula(cell.value):
                        continue
                    computed = counterpart[cell.coordinate]
                    result = computed.value
                    error_value = result if computed.data_type == 'e' else None
                    if error_value:
                        errors += 1
                    rows.append({
                        'sheet': sheet.title,
                        'cell': cell.coordinate,
                        **formula_record(cell.value),
                        'engine_value': None if error_value else normalize_value(result),
                        'engine_error': error_value,
                    })
        delivery = None
        if destination is not None:
            destination.parent.mkdir(parents=True, exist_ok=True)
            # Copy the native engine document. Saving a data_only workbook erases formulas.
            with tempfile.TemporaryDirectory(prefix='.workbook-recalculate-', dir=destination.parent) as staging:
                staged = Path(staging) / 'result.xlsx'
                shutil.copyfile(converted, staged)
                os.replace(staged, destination)
            delivery = str(destination)
    return {
        'source': str(source),
        'delivery': delivery,
        'engine': detail,
        'formula_count': len(rows),
        'error_value_count': errors,
        'values_are_engine_results': True,
        'note': 'engine_value/engine_error are outputs produced by the installed '
                'spreadsheet engine from the source bytes; null means the engine '
                'produced no value for that cell. Each row represents a stored formula anchor; '
                'array/data-table ranges and inputs are retained as formula_attributes. '
                'Data tables have no scalar formula text, so formula is null.',
        'formulas': rows,
    }


OPERATIONS = {
    'inspect': operation_inspect,
    'write': operation_write,
    'recalculate': operation_recalculate,
}


def main():
    args = json.load(sys.stdin)
    operation = OPERATIONS.get(args.get('operation'))
    if operation is None:
        raise ValueError('Unknown workbook operation')
    emit(operation(args))


if __name__ == '__main__':
    try:
        main()
    except Exception as error:  # validation and engine failures must propagate
        sys.stderr.write('%s: %s\n' % (type(error).__name__, error))
        sys.exit(1)
