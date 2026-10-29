import importlib.util
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment

spec = importlib.util.spec_from_file_location('candidate', '/candidate/workbook.py')
candidate = importlib.util.module_from_spec(spec)
spec.loader.exec_module(candidate)

book = Workbook()
sheet = book.active
sheet.title = 'Scenarios'
sheet.append(['Metric', 'Case A', 'Case B', 'Case C'])
sheet.append(['Exit enterprise value', 120, 150, 180])
exit_findings = candidate.exit_row_findings(
    {2: {'label': sheet['A2'].value, 'role': candidate.role(sheet['A2'].value)}},
    {2: {2: 120}, 3: {2: 150}, 4: {2: 180}}, [2, 3, 4],
)
assert exit_findings, 'Expected to reproduce scenario-columns misclassification'

sheet['A3'] = 'Gross margin (percentage points)'
sheet['B3'] = 25
sheet['B3'].number_format = 'General'
margin_findings = candidate.margin_scale_findings(
    sheet, {3: {'label': sheet['A3'].value, 'role': 'margin'}}, {2: {3: 25}}, [2],
)
assert margin_findings, 'Expected to reproduce unrequested scale inference'

formula_book = Workbook()
formula_sheet = formula_book.active
formula_sheet['A1'] = '="#N/A"'
cache_book = Workbook()
cache_book.active['A1'] = '#N/A'
cache_book.active['A1'].data_type = 's'
error_findings = candidate.write_back_findings(
    formula_book, formula_sheet, cache_book, formula_sheet.title,
)
assert error_findings, 'Expected to reproduce literal-text error misclassification'

sheet['A2'].alignment = Alignment(indent=2)
reported_indent = candidate.indent_of(book, sheet['A2'])
assert reported_indent == 0 and sheet['A2'].alignment.indent == 2

print(json.dumps({
    'candidate': '01a983e0ba302871c2b1836942f3f07362b0f246db33994d74e991420b6bcfa4',
    'boundary': 'Deterministic direct helper counterexamples in the frozen runtime; no model calls, no evaluator changes, no changes to recorded research scores.',
    'cases': [
        {'name': 'parallel-scenarios', 'input': 'Three parallel exit scenarios, each legitimately populated.', 'observation': exit_findings},
        {'name': 'explicit-percentage-points', 'input': 'Gross margin explicitly expressed in percentage points; 25 means 25%.', 'observation': margin_findings},
        {'name': 'literal-error-looking-text', 'input': 'Formula returns the ordinary string #N/A; cached cell data_type is s.', 'observation': error_findings},
        {'name': 'alignment-indent', 'input': {'alignmentIndent': 2}, 'observation': {'reportedIndent': reported_indent}},
    ],
}, indent=2))
